"""Download -> clean -> load. Run: python -m etl [--file PATH] [--inspect] [--dry-run] [--force]"""
import argparse
import hashlib
import json
import sys

import openpyxl
import pandas as pd

from .clean import clean_sheet, column_kind, detect_header_row, build_headers
from .config import load_settings
from .items import check_mapping, load_catalog_groups, load_mapping, to_items
from .load import Supabase
from .source import fetch


def content_hash(records: list[dict]) -> str:
    return hashlib.sha256(json.dumps(records, sort_keys=True, ensure_ascii=False, default=str).encode("utf-8")).hexdigest()


def read_sheet(ws) -> pd.DataFrame:
    # Index = Excel row number - 1, so source rows stay traceable after cleaning.
    return pd.DataFrame(list(ws.iter_rows(min_row=1, values_only=True)), dtype=object)


def inspect(wb) -> None:
    for ws in wb.worksheets:
        raw = read_sheet(ws).dropna(how="all").dropna(axis=1, how="all")
        print(f"\n=== Sheet '{ws.title}' ({ws.sheet_state}) - {len(raw)} dòng có dữ liệu")
        if raw.empty:
            continue
        pos = detect_header_row(raw)
        print(f"Dòng tiêu đề dò được: {int(raw.index[pos]) + 1}")
        sample = raw.iloc[pos + 1 : pos + 4]
        for i, (orig, name) in enumerate(zip(raw.iloc[pos], build_headers(list(raw.iloc[pos])))):
            values = [str(v)[:30] for v in sample.iloc[:, i] if v is not None and v == v]
            print(f"  {name:<35} [{column_kind(name):<6}] <- {str(orig)[:40]!r}  vd: {values}")


def main() -> int:
    # Windows consoles default to cp1252, which cannot print Vietnamese.
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap =argparse.ArgumentParser(prog="etl")
    ap.add_argument("--file", help="Đọc file local thay vì tải từ SharePoint")
    ap.add_argument("--inspect", action="store_true", help="Chỉ in cấu trúc sheet/cột rồi thoát")
    ap.add_argument("--dry-run", action="store_true", help="Làm sạch và ghi ra out/, không đẩy lên Supabase")
    ap.add_argument("--force", action="store_true", help="Đẩy lên kể cả khi nội dung sheet không đổi")
    args = ap.parse_args()

    s = load_settings()
    src = fetch(s, args.file)
    print(f"Nguồn: {src.path} (eTag={src.etag}, sửa lúc {src.modified_at})")

    wb = openpyxl.load_workbook(src.path, read_only=True, data_only=True)
    if args.inspect:
        inspect(wb)
        return 0

    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    sheets = s.sheets or visible[:1]
    missing = [name for name in sheets if name not in wb.sheetnames]
    if missing:
        print(f"Không thấy sheet: {missing}. Có: {wb.sheetnames}", file=sys.stderr)
        return 1

    if not s.product_catalog_file:
        print("Thiếu PRODUCT_CATALOG_FILE trong .env", file=sys.stderr)
        return 1
    mapping, groups = load_mapping(), load_catalog_groups(s.product_catalog_file)
    unknown = check_mapping(mapping, groups)
    if unknown:
        print(f"Mapping trỏ tới nhóm không có trong danh mục sản phẩm: {unknown}", file=sys.stderr)
        return 1
    results = [to_items(clean_sheet(read_sheet(wb[name]), name, s.header_row), mapping, groups) for name in sheets]

    s.out_dir.mkdir(parents=True, exist_ok=True)
    all_issues = []
    for r in results:
        pd.DataFrame(r.records).to_csv(s.out_dir / f"cleaned_{r.sheet}.csv", index=False, encoding="utf-8-sig")
        all_issues += r.issues
        print(
            f"[{r.sheet}] tiêu đề dòng {r.header_row}, {len(r.records)} dòng mặt hàng, "
            f"bỏ {r.skipped_rows} dòng gói/lô, {len(r.issues)} lỗi"
        )
    pd.DataFrame(all_issues, columns=["sheet", "excel_row", "column", "kind", "value"]).to_csv(
        s.out_dir / "issues.csv", index=False, encoding="utf-8-sig"
    )
    print(f"Đã ghi kết quả làm sạch vào {s.out_dir}")

    if args.dry_run:
        return 0

    sb = Supabase(s)
    loaded = sb.sheet_hashes()
    hashes = {r.sheet: content_hash(r.records) for r in results}
    changed = [r for r in results if args.force or loaded.get(r.sheet) != hashes[r.sheet]]
    for r in results:
        if all(r is not c for c in changed):
            print(f"[{r.sheet}] không đổi so với lần nạp trước, bỏ qua.")
    if not changed:
        return 0

    run_id = sb.start_run(src.name, src.etag, src.modified_at)
    try:
        total = sum(sb.replace_sheet(run_id, r.sheet, hashes[r.sheet], r.records) for r in changed)
    except Exception as exc:
        sb.finish_run(run_id, "failed", error=str(exc)[:2000])
        raise
    sb.finish_run(run_id, "success", row_count=total, issue_count=len(all_issues))
    print(f"Sync #{run_id} xong: {total} dòng lên Supabase.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
